import tkinter as tk
from tkinter import ttk
from tkinter import messagebox
from tkinter import simpledialog
import openpyxl
from datetime import datetime, date
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg

def ahorros_list():
        filepath = "./Gastos.xlsx"
        workbook = openpyxl.load_workbook(filepath)
        sheet = workbook["Ahorros"]
        ahorros = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[0] not in ahorros:
                ahorros.append(row[0])
        return ahorros

def new_ahorro(ahorro_select, ahorro_amount, ahorro_final, charts_frame, canvas_scroll):
    
    filepath = "./Gastos.xlsx"
    workbook = openpyxl.load_workbook(filepath)
    sheet = workbook["Ahorros"]
    ahorro_name = ahorro_select
    
    try:
        amount_added = float(ahorro_amount)
        objective = float(ahorro_final)
        
        # Busca el ahorro existente y actualiza el monto total
        updated = False
        for row in sheet.iter_rows(min_row=2):
            if row[0].value == ahorro_name:
                # Suma el monto agregado al total guardado
                total_amount = row[1].value + amount_added
                sheet.cell(row=row[0].row, column=2, value=total_amount)  # Actualiza el monto total en la hoja
                sheet.cell(row=row[0].row, column=5, value=amount_added)  # Registra el monto agregado en la última columna
                updated = True
                break
            
        # Si no existe un registro previo, crea uno nuevo
        if not updated:
            sheet.append([ahorro_name, amount_added, objective, date.today(), amount_added])
        
    except ValueError:
        messagebox.showerror("Error", "Debe ingresar valores numéricos válidos.")
        return
    
    workbook.save(filepath)
    
    # Actualiza los gráficos
    update_charts(charts_frame, canvas_scroll)


    
def get_values():
    filepath = "./Gastos.xlsx"
    workbook = openpyxl.load_workbook(filepath)
    sheet = workbook["Ahorros"]
    if sheet.max_row == 1:  # Check if the sheet is empty (only header row)
        
        return
    ahorros = []
    amounts_saved = []
    objectives = []
    dates = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] not in ahorros:
            ahorros.append(row[0])
            amounts_saved.append(row[1])
            objectives.append(row[2])
            if isinstance(row[3], str):
                try:
                    dates.append(datetime.strptime(row[3], "%Y-%m-%d"))
                except ValueError:
                    dates.append(datetime.strptime(row[3], "%d/%m/%Y"))
            else:
                dates.append(row[3])
    sorted_data = sorted(zip(ahorros, amounts_saved, objectives, dates), key=lambda x: x[3], reverse=True)
    latest_data = []
    for ahorro, _, _, _ in sorted_data:
        if ahorro not in [x[0] for x in latest_data]:
            latest_data.append(next((x for x in sorted_data if x[0] == ahorro), None))
    ahorros, amounts_saved, objectives, dates = zip(*latest_data)
    return ahorros, amounts_saved, objectives, dates

def clear_charts(charts_frame):
    for widget in charts_frame.winfo_children():
        widget.destroy()

def update_charts(charts_frame,canvas_scroll):
    filepath = "./Gastos.xlsx"
    workbook = openpyxl.load_workbook(filepath)
    sheet = workbook["Ahorros"]
    if sheet.max_row == 1:  # Check if the sheet is empty (only header row)
        
        return    
    clear_charts(charts_frame)
    ahorros, amounts_saved, objectives, dates = get_values()
    fig, axs = plt.subplots(len(ahorros), 1, figsize=(12, 1.1 * len(ahorros)))
    fig.suptitle('Listado de Ahorros', color="#bfbfbf")
    
    if len(ahorros) == 1:
        axs = [axs]
    
    for i, (ahorro, amount_saved, objective) in enumerate(zip(ahorros, amounts_saved, objectives)):
        axs[i].barh(0, amount_saved, height=0.3, color='#84b86d')
        axs[i].get_yaxis().set_visible(False)
        axs[i].spines['top'].set_visible(False)
        axs[i].spines['right'].set_visible(False)
        axs[i].spines['left'].set_visible(False)
        axs[i].spines['bottom'].set_color('#808080')
        axs[i].tick_params(axis='x', colors='#808080')
        axs[i].set_yticks([0])
        axs[i].set_yticklabels([ahorro])
        axs[i].set_facecolor('#313131')
        axs[i].set_ylim(-0.4, 0.8)
        axs[i].set_xlim(0, objective)
        axs[i].set_xticks([0, objective])
        axs[i].set_title(ahorro, fontsize=10, loc='left', fontstyle='italic', fontweight='book',
                         bbox=dict(facecolor='#b6d7a8', edgecolor='black', boxstyle='round,pad=0.5'))
        if objective-amount_saved > 0:
            axs[i].text(0.95, 0.98, f"Restante: ${int(objective - amount_saved):,}", ha="right", va="center",
                    fontweight='bold', color='black', fontsize=9,
                    bbox=dict(facecolor='#bac4c1', edgecolor='black', boxstyle='round,pad=0.5'),
                    transform=axs[i].transAxes)
        else:
            axs[i].text(0.95, 0.98, f"Excedente: ${int(amount_saved-objective):,}", ha="right", va="center",
                    fontweight='bold', color='#309651', fontsize=9,
                    bbox=dict(facecolor='#bac4c1', edgecolor='black', boxstyle='round,pad=0.5'),
                    transform=axs[i].transAxes)
    fig.tight_layout()
    fig.patch.set_facecolor("#313131")
    canvas = FigureCanvasTkAgg(fig, master=charts_frame)
    canvas.draw()
    canvas.get_tk_widget().pack(side=tk.TOP, fill=tk.BOTH, expand=1)
    # Update scrollbar region
    charts_frame.update_idletasks()
    canvas_scroll.config(scrollregion=canvas_scroll.bbox("all"))

def insert_table(selected_item,treeview):
    selected_item
    for item in treeview.get_children():
        treeview.delete(item)
    filepath = "./Gastos.xlsx"
    workbook = openpyxl.load_workbook(filepath)
    sheet = workbook["Ahorros"]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] == selected_item:
            treeview.insert("", "end", values=(row[3], row[1], row[2]))