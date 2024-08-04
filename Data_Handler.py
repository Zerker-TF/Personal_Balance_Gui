# Here we make all the magic numbers calculations
# We get the values inserted by the user and store them to return where are needed
import openpyxl
import os
import datetime

# Show_chart data for both charts
def category_totals():
   categorias =["Comida","Alquiler","Expensas","Internet","Agua","Gas","Luz","Transporte","Salud","Bolucompra"]
   filepath = "./Gastos.xlsx"
   workbook = openpyxl.load_workbook(filepath)
   selected_month = month_select.get()
   #print(selected_month)
   try:
     sheet = workbook[selected_month]
   except KeyError:
     sheet = workbook.active
   
   amounts = [sum(float(row[2]) for row in data if row[1] == category) for category in categorias]
   sorted_categories = [category for _, category in sorted(zip(amounts, categorias), reverse=True)]
   sorted_amounts = [amount for amount, _ in sorted(zip(amounts, categorias), reverse=True)]
   total = sum(amounts)
   
   # Percentages for pie chart   
   percentages = [f"{(amount/total)*100:.2f}%" for amount in amounts]
   
    # Create a new category "Otros" for pie chart
   otros_amount = sum(amount for amount, percentage in zip(amounts, percentages) if float(percentage.strip('%')) < 3)
   otros_percentage = f"{(otros_amount/total)*100:.2f}%"

def calcular_saldo(filepath, sheetname):
    datos = obtener_datos(filepath, sheetname)
    total_ingresos = sum(float(row[2]) for row in datos if row[1] == "Ingresos")
    total_gastos = sum(float(row[2]) for row in datos if row[1] != "Ingresos" and row[1] != "Ahorros")
    saldo = total_ingresos - total_gastos
    return saldo

     
def datos():
   filepath = "./Gastos.xlsx"
   # Check if excel file exist in path
   if not os.path.exists(filepath):
      workbook = openpyxl.Workbook()
      months = ["Enero","Febrero","Marzo","Abril","Mayo","Junio","Julio","Agosto","Septiembre","Octubre","Noviembre","Diciembre"]
      for month_name in months:
         sheet = workbook.create_sheet(title=month_name)
         heading = ["Fecha","Categoria","Monto","Descripcion","Cuotas"]
         sheet.append(heading)
      workbook.save(filepath)
      
   workbook = openpyxl.load_workbook(filepath)
   #print(month)
   try:
      sheet = workbook[month]
   except KeyError:
      sheet = workbook.active
   
   #list_values = list(sheet.values)
   list_values = [list(row) for row in sheet.values]
   # Sort the list by date
   for i, row in enumerate(list_values[1:]):
      if isinstance(row[0], str):
         list_values[i+1][0] = datetime.datetime.strptime(row[0], "%Y-%m-%d").date()
      elif isinstance (row[0], datetime.datetime):
         list_values[i+1][0] =row[0].date()
         
   list_values[1:] = sorted(list_values[1:], key=lambda x: x[0])