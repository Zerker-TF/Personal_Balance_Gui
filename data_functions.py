import tkinter as tk
from tkinter import ttk
import os
import openpyxl, openpyxl.workbook
from tkcalendar import Calendar, DateEntry
import datetime


def load_data(month, treeview):
    
    months = ["Enero","Febrero","Marzo","Abril","Mayo","Junio","Julio","Agosto","Septiembre","Octubre","Noviembre","Diciembre"]
   
    filepath = "./Gastos.xlsx"
    
    if not os.path.exists(filepath):
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        
        ahorros_sheet = workbook.create_sheet(title="Ahorros",index=0)
        ahorros_heading = ["Fecha","Nombre","Total","Objetivo","Monto agregado"]
        ahorros_sheet.append(ahorros_heading)
        
        
        for name in months:
            sheet = workbook.create_sheet(title=name)
            heading = ["Fecha","Categoria","Monto","Descripcion","Cuotas"]
            sheet.append(heading)
        workbook.save(filepath)
    
    workbook = openpyxl.load_workbook(filepath)
    
    try:
        sheet = workbook[month]
    except KeyError:
        sheet = workbook.active
    
    list_values = [list(row) for row in sheet.values]
    #Sort by date
    for i, row in enumerate(list_values[1:]):
        if isinstance(row[0],str):
            list_values[i+1][0] = datetime.datetime.strptime(row[0], "%Y-%m-%d").date()
        elif isinstance (row[0], datetime.datetime):
            list_values[i+1][0] = row[0].date()
    
    list_values[1:] = sorted(list_values[1:], key=lambda x: x[0])
    for item in treeview.get_children():
        treeview.delete(item)
    for row in list_values[1:]:
        treeview.insert("","end",values=row[0:])
        
    pass

def insert_data(exp_date,treeview):
    months = ["Enero","Febrero","Marzo","Abril","Mayo","Junio","Julio","Agosto","Septiembre","Octubre","Noviembre","Diciembre"]
    date = datetime.datetime.strptime(exp_date, "%m/%d/%y")
    month = int(date.strftime("%m"))
    print(month)
    month_str = months[month-1]
    print(month_str)
    
    pass