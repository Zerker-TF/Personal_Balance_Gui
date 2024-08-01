import tkinter as tk
from tkinter import ttk
import openpyxl
import os
import openpyxl.workbook
from tkcalendar import Calendar, DateEntry
import datetime
import gui
from graphs import show_chart

def monthly_total(month, amount, category):
 
   global category_totals, months
   
   month_index = months.index(month)
   
   category_index = gui.combo_list.index(category)
   
   category_totals[month_index][category_index] += amount
   
   return category_totals[month_index]

# Save insterted data into the correct excel sheet
def insert_row(expense_date,expense_type,expense_value,expense_entry,expense_cuota):
   date = expense_date.get_date().strftime("%Y-%m-%d")
   month = date[5:7]
   year = int(date[:4])
      
   category = expense_type.get()
   amount = float(expense_value.get())
   cuota = expense_cuota.get()
   desc = expense_entry.get()
   # saving the total of each category ("Comida","Alquiler","Expensas","Internet","Agua","Gas","Luz","Transporte","Salud","Bolucompra","Ingresos","Ahorros")
   #print(months[int(month)-1])
   
          
   filepath = "./Gastos.xlsx"
   if not os.path.exists(filepath):
      workbook = openpyxl.Workbook()
      months = ["Enero","Febrero","Marzo","Abril","Mayo","Junio","Julio","Agosto","Septiembre","Octubre","Noviembre","Diciembre"]
   for month_name in months:
    sheet = workbook.create_sheet(title=month_name)
    heading = ["Fecha","Categoria","Monto","Descripcion","Cuotas"]
    sheet.append(heading)
    workbook.save(filepath)
   
   monthly_total(months[int(month)-1], amount, category)   
   
   workbook = openpyxl.load_workbook(filepath)
   
   #try:
   #   sheet = workbook[months[int(month)-1]]
   #except KeyError:
   #   sheet = workbook.active
   
   # calculate cuotas
   if cuota != "1":
      cuota = int(cuota)
      divided_amount = amount / cuota
      month_index = int(month)
      
      for i in range(cuota):
         new_date = f"{year}-{str(month_index).zfill(2)}-{date[8:]}"
         new_row = [new_date, category, divided_amount, desc if desc else " ", cuota]
         
         sheet = workbook[months[month_index-1]]         
         sheet.append(new_row)
         # Move to the next month
         month_index += 1
         if month_index > 12:
            month_index = 1
            year += 1
         
   else:
      sheet = workbook[months[int(month)-1]]
      new_row = [date, category, amount, desc if desc else " ", cuota if cuota else 1]
      sheet.append(new_row)     
      
   workbook.save(filepath)
   
   gui.treeview.insert("","end",values=[date,category,amount,desc if desc else " "])
   if not gui.chart_shown:
      show_chart(True)

# Load data from excel file on the selected month
def load_data():
   month = gui.month_select.get()

   filepath = "./Gastos.xlsx"
   # Check if excel file exist in path
   if not os.path.exists(filepath):
      workbook = openpyxl.Workbook()
      months = ["Enero","Febrero","Marzo","Abril","Mayo","Junio","Julio","Agosto","Septiembre","Octubre","Noviembre","Diciembre"]
      for month_name in months:
         sheet = workbook.create_sheet(title=month_name)
         heading = ["Fecha","Categoria","Monto","Descripcion","Cuotas"]
         sheet.append(heading)
      workbook.save(filepath)
      
   workbook = openpyxl.load_workbook(filepath)
   #print(month)
   try:
      sheet = workbook[month]
   except KeyError:
      sheet = workbook.active
   
   #list_values = list(sheet.values)
   list_values = [list(row) for row in sheet.values]
   # Sort the list by date
   for i, row in enumerate(list_values[1:]):
      if isinstance(row[0], str):
         list_values[i+1][0] = datetime.datetime.strptime(row[0], "%Y-%m-%d").date()
      elif isinstance (row[0], datetime.datetime):
         list_values[i+1][0] =row[0].date()
         
   list_values[1:] = sorted(list_values[1:], key=lambda x: x[0])
   #clear the treeview window before printing
   for item in gui.treeview.get_children():
      gui.treeview.delete(item)
      
   for row in list_values[1:]:
      gui.treeview.insert("","end",values=row[0:])
   if not gui.chart_shown:
      show_chart(True)