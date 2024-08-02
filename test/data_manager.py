import tkinter as tk
from tkinter import ttk
import openpyxl
import os
import openpyxl.workbook
from tkcalendar import Calendar, DateEntry
import datetime
from tkinter import messagebox
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
import pyperclip



def show_chart(graphframe,frame,root,month_select,chart_shown,refresh=False):
    global  canvas, pie_canvas
    colors = {

     "Comida": ['#4CAF50', '#2E865F'],
 
     "Alquiler": ['#FF9800', '#FFC107'],
 
     "Expensas": ['#009688', '#00796B'],
 
     "Internet": ['#2196F3', '#1976D2'],
 
     "Agua": ['#66c0f4', '#45B3FA'],
 
     "Gas": ['#c7d5e0', '#B2E6CE'],
 
     "Luz": ['#ffd700', '#F7DC6F'],
 
     "Transporte": ['#ff7f50', '#cc6540'],
 
     "Salud": ['#3E69DA', '#2b4998'],
 
     "Bolucompra": ['#d11141', '#a70d34'],

      }
    if chart_shown or refresh:

        # Set the active sheet based on the selected month
        filepath = "./Gastos.xlsx"
        workbook = openpyxl.load_workbook(filepath)
        selected_month = month_select.get()
        #print(selected_month)
        try:
          sheet = workbook[selected_month]
        except KeyError:
          sheet = workbook.active
    
        data = list(sheet.values)
        categorias =["Comida","Alquiler","Expensas","Internet","Agua","Gas","Luz","Transporte","Salud","Bolucompra"]

        amounts = [sum(float(row[2]) for row in data if row[1] == category) for category in categorias]
        sorted_categories = [category for _, category in sorted(zip(amounts, categorias), reverse=True)]
        sorted_amounts = [amount for amount, _ in sorted(zip(amounts, categorias), reverse=True)]
        total = sum(amounts)

        percentages = [f"{(amount/total)*100:.2f}%" for amount in amounts]

        # Create a new category "Otros" for pie chart
        otros_amount = sum(amount for amount, percentage in zip(amounts, percentages) if float(percentage.strip('%')) < 3)
        otros_percentage = f"{(otros_amount/total)*100:.2f}%"

        if otros_amount > 0:
          pie_categorias = [category for category, percentage in zip(categorias, percentages) if float(percentage.strip('%')) >= 3] + ["Otros"]
          pie_amounts = [amount for amount, percentage in zip(amounts, percentages) if float(percentage.strip('%')) >= 3] + [otros_amount]
        else:
          pie_categorias = [category for category, percentage in zip(categorias, percentages) if float(percentage.strip('%')) >= 3]
          pie_amounts = [amount for amount, percentage in zip(amounts, percentages) if float(percentage.strip('%')) >= 3] 


        # Create the figure and axis
        figure, ax = plt.subplots(figsize=(8, 6))
        # Create the figure and axis for pie chart
        figure2, ax2 = plt.subplots(figsize=(4.5, 6))

        # Title and labels
        figure.suptitle(f"Total gastado por categoria para el mes {selected_month}", y=0.98)
        ax.set_xlabel("Categoria")
        ax.set_ylabel("Total ($ARS)")

        # Plot the bar chart  
        bars = ax.bar(sorted_categories, sorted_amounts, color=[colors[category][0] for category in sorted_categories], edgecolor='black', linewidth=1)
        # Add shadows
        ax.bar([x - 0.5 for x in range(len(sorted_categories))], sorted_amounts, color=[colors[category][1] for category in sorted_categories], edgecolor='black', linewidth=1, zorder=10)

        ax.set_xticks([x - 0.25 for x in range(len(categorias))])
        ax.set_xticklabels(sorted_categories, rotation=45, ha='right')
        ax.set_facecolor('#c0c0c0')
        ax.text(0.95, 0.95, f"Total: ${int(total):,}", ha="right", va="top", 
           fontweight='bold', color='black', fontsize=10, transform=ax.transAxes, bbox=dict(facecolor='white', edgecolor='black', boxstyle='round,pad=0.5'))
        figure.patch.set_facecolor("#808080")
        figure.tight_layout() # makes the labels fit the plot area

        # Create the canvas for the bar chart
        canvas = FigureCanvasTkAgg(figure, master=graphframe)
        canvas.draw()
        canvas.get_tk_widget().grid(row=0, column=0)

        # pctdistance moves the label x amount from the center to the outside
        ax2.pie(pie_amounts,  autopct=lambda p: '{:.1f}%'.format(p), startangle=90, pctdistance=0.75, textprops={'fontsize':10, 'fontweight': 'bold'}  ,colors=['#4CAF50', '#FF9800', '#009688', '#2196F3', '#66c0f4', '#c7d5e0', '#ffd700', '#ff7f50', '#d11141', '#808080'], labels=pie_categorias)
        ax2.axis('equal')
        figure2.suptitle("Porcentaje por categoria", y=0.98)
        figure2.patch.set_facecolor("#808080")

        # Add each category total to the top of each bar
        for bar, amount in zip(bars.patches, sorted_amounts):
           if int(amount) >= 1000:
              text = f"{amount/1000:.1f}k"
           else:
              text= f"{int(amount):,}"
           ax.text(bar.get_x() - 0.15 + bar.get_width()/2, bar.get_y() + bar.get_height(), text, ha='center', va='bottom', fontweight='bold', color='black',fontsize=9)

        # Canvas for pie chart
        pie_canvas = FigureCanvasTkAgg(figure2, master=graphframe)
        pie_canvas.draw()
        pie_canvas.get_tk_widget().grid(row=0, column=1)


        # Resize window when chart is shown
        graphframe.grid_rowconfigure(0, weight=1)
        graphframe.grid_columnconfigure(0, weight=1)
        frame.rowconfigure(graphframe, weight=1)
        frame.columnconfigure(graphframe, weight=1)
        frame.rowconfigure(1, weight=1)
        frame.columnconfigure(0, weight=1)
        root.geometry("1100x780")
        chart_shown = False
    else:
       # Hide the chart
       for widget in graphframe.winfo_children():
         widget.grid_remove()
       graphframe.grid_rowconfigure(0, weight=0)
       graphframe.grid_columnconfigure(0, weight=0)
       frame.rowconfigure(1, weight=0)
       frame.columnconfigure(0, weight=0)
       root.geometry("1100x356")
       chart_shown = True
    return chart_shown
   
def insert_row(expense_date,expense_type,expense_value,expense_cuota,expense_entry,treeview):
   date = expense_date.get_date().strftime("%Y-%m-%d")
   month = date[5:7]
   year = int(date[:4])
      
   category = expense_type.get()
   amount = float(expense_value.get())
   cuota = expense_cuota.get()
   desc = expense_entry.get()
   chart_shown = True
          
   filepath = "./Gastos.xlsx"
   if not os.path.exists(filepath):
      workbook = openpyxl.Workbook()
      months = ["Enero","Febrero","Marzo","Abril","Mayo","Junio","Julio","Agosto","Septiembre","Octubre","Noviembre","Diciembre"]
   for month_name in months:
    sheet = workbook.create_sheet(title=month_name)
    heading = ["Fecha","Categoria","Monto","Descripcion","Cuotas"]
    sheet.append(heading)
    workbook.save(filepath)
   
   monthly_total(months[int(month)-1], amount, category)   
   
   workbook = openpyxl.load_workbook(filepath)
   
   
   if cuota != "1":
      cuota = int(cuota)
      divided_amount = amount / cuota
      month_index = int(month)
      
      for i in range(cuota):
         new_date = f"{year}-{str(month_index).zfill(2)}-{date[8:]}"
         new_row = [new_date, category, divided_amount, desc if desc else " ", cuota]
         
         sheet = workbook[months[month_index-1]]         
         sheet.append(new_row)
         # Move to the next month
         month_index += 1
         if month_index > 12:
            month_index = 1
            year += 1
         
   else:
      sheet = workbook[months[int(month)-1]]
      new_row = [date, category, amount, desc if desc else " ", cuota if cuota else 1]
      sheet.append(new_row)     
      
   workbook.save(filepath)
   
   treeview.insert("","end",values=[date,category,amount,desc if desc else " "])
   if not chart_shown:
      show_chart(True)
 
def monthly_total(month, amount, category):
    global category_totals
    months = ["Enero","Febrero","Marzo","Abril","Mayo","Junio","Julio","Agosto","Septiembre","Octubre","Noviembre","Diciembre"]
    combo_list = ["Comida","Alquiler","Expensas","Internet","Agua","Gas","Luz","Transporte","Salud","Bolucompra","Ingresos","Ahorros"]
    
    month_index = months.index(month)
    
    category_index = combo_list.index(category)
    
    category_totals[month_index][category_index] += amount
    
    return category_totals[month_index]    

def load_data(month_select, treeview, chart_shown,graph,frame,root):
   month = month_select.get()

   filepath = "./Gastos.xlsx"
   # Check if excel file exist in path
   if not os.path.exists(filepath):
      workbook = openpyxl.Workbook()
      months = ["Enero","Febrero","Marzo","Abril","Mayo","Junio","Julio","Agosto","Septiembre","Octubre","Noviembre","Diciembre"]
      for month_name in months:
         sheet = workbook.create_sheet(title=month_name)
         heading = ["Fecha","Categoria","Monto","Descripcion","Cuotas"]
         sheet.append(heading)
      workbook.save(filepath)
      
   workbook = openpyxl.load_workbook(filepath)
   #print(month)
   try:
      sheet = workbook[month]
   except KeyError:
      sheet = workbook.active
   
   #list_values = list(sheet.values)
   list_values = [list(row) for row in sheet.values]
   # Sort the list by date
   for i, row in enumerate(list_values[1:]):
      if isinstance(row[0], str):
         list_values[i+1][0] = datetime.datetime.strptime(row[0], "%Y-%m-%d").date()
      elif isinstance (row[0], datetime.datetime):
         list_values[i+1][0] =row[0].date()
         
   list_values[1:] = sorted(list_values[1:], key=lambda x: x[0])
   #clear the treeview window before printing
   for item in treeview.get_children():
      treeview.delete(item)
      
   for row in list_values[1:]:
      treeview.insert("","end",values=row[0:])
   if not chart_shown:
      show_chart(graph, frame, root, month_select, chart_shown, refresh=True)
      chart_shown = True  # Set chart_shown to True after showing the chart
   else:
      chart_shown = False 
   return chart_shown
#def copy_row():
#    #
#def delete_row():
#    #
#def edit_row():
#    #
#def show_menu(event):
#    #
#def close_menu(event):