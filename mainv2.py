import tkinter as tk
from tkinter import ttk
from Data_Expensesv2 import create_expense_window
from Data_Savingv2 import create_savings_window
import openpyxl
import os
def resize(event):
    selected_tab = event.widget.select()
    tab_text = event.widget.tab(selected_tab, "text")

    if tab_text == "Gastos":
        root.minsize(1260, 410)
        root.maxsize(1260, 950)
       
    elif tab_text == "Ahorros":
        root.minsize(1260, 880)
        root.maxsize(1260, 945)
        root.wm_geometry("1260x900")

def initialize_sheets():
    filepath = "./Gastos.xlsx"
    months = ["Enero","Febrero","Marzo","Abril","Mayo","Junio","Julio","Agosto","Septiembre","Octubre","Noviembre","Diciembre"]
    if not os.path.exists(filepath):
     workbook = openpyxl.Workbook()
     workbook.remove(workbook.active)

     ahorros_sheet = workbook.create_sheet(title="Ahorros",index=0)
     ahorros_heading = ["Fecha","Nombre","Total","Objetivo","Monto agregado"]
     ahorros_sheet.append(ahorros_heading)


     for name in months:
         sheet = workbook.create_sheet(title=name)
         heading = ["Fecha","Categoria","Monto","Descripcion","Cuotas"]
         sheet.append(heading)
     workbook.save(filepath)
    
def main():
    global root
    initialize_sheets()
    root = tk.Tk()
    root.title("Balance Personal")
    root.minsize(1260,400)
    root.maxsize(1260,400)
    root.grid_rowconfigure(0,weight=1)
    root.columnconfigure(0,weight=1)
    
    Ventana = ttk.Notebook(root)
    Ventana.grid(row=0,column=0,sticky="new")
    Ventana.columnconfigure(0,weight=1)
    
    
    frame_gastos = ttk.Frame(Ventana)
    frame_ahorros =ttk.Frame(Ventana)
    
    Ventana.add(frame_gastos,text="Gastos")
    Ventana.add(frame_ahorros,text="Ahorros")
    
    create_expense_window(frame_gastos)
    create_savings_window(frame_ahorros)
    
    Ventana.bind("<<NotebookTabChanged>>", resize)
    
    root.mainloop()
    
if __name__ == "__main__":
    main()
    